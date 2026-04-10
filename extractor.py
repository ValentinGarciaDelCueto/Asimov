# ============================================================
#  extractor.py — Extrae texto de PDF y DOCX
# ============================================================

import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_text(file_path: Path) -> str:
    """
    Extrae texto de un archivo PDF o DOCX.
    Retorna el texto completo como string, o "" si falla.
    """
    ext = file_path.suffix.lower()

    if ext == ".pdf":
        return _extract_pdf(file_path)
    elif ext == ".docx":
        return _extract_docx(file_path)
    else:
        logger.warning(f"Extensión no soportada: {ext} — {file_path.name}")
        return ""


def _extract_pdf(file_path: Path) -> str:
    """Extrae texto de un PDF usando pdfplumber."""
    try:
        import pdfplumber

        text_parts = []
        with pdfplumber.open(file_path) as pdf:
            total_pages = len(pdf.pages)
            logger.info(f"  PDF: {total_pages} páginas — {file_path.name}")

            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
                if (i + 1) % 10 == 0:
                    logger.info(f"  Procesando página {i + 1}/{total_pages}...")

        full_text = "\n".join(text_parts).strip()

        if not full_text:
            logger.warning(f"  No se extrajo texto de {file_path.name} (¿PDF escaneado?)")

        return full_text

    except ImportError:
        logger.error("pdfplumber no instalado. Ejecutá: pip install pdfplumber")
        return ""
    except Exception as e:
        logger.error(f"Error al leer PDF {file_path.name}: {e}")
        return ""


def _extract_docx(file_path: Path) -> str:
    """Extrae texto de un archivo Word (.docx)."""
    try:
        from docx import Document

        doc = Document(str(file_path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # También extrae texto de tablas
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())

        full_text = "\n".join(paragraphs).strip()

        logger.info(f"  DOCX: {len(paragraphs)} párrafos — {file_path.name}")

        if not full_text:
            logger.warning(f"  No se extrajo texto de {file_path.name}")

        return full_text

    except ImportError:
        logger.error("python-docx no instalado. Ejecutá: pip install python-docx")
        return ""
    except Exception as e:
        logger.error(f"Error al leer DOCX {file_path.name}: {e}")
        return ""


def split_into_chunks(text: str, chunk_size: int = 3000) -> list[str]:
    """
    Divide el texto en chunks de N palabras con superposición
    para no perder contexto entre fragmentos.
    """
    words = text.split()
    if len(words) <= chunk_size:
        return [text]

    chunks = []
    overlap = chunk_size // 10  # 10% de superposición
    step = chunk_size - overlap
    i = 0

    while i < len(words):
        chunk_words = words[i : i + chunk_size]
        chunks.append(" ".join(chunk_words))
        i += step

    logger.info(f"  Texto dividido en {len(chunks)} chunks")
    return chunks


def count_words(text: str) -> int:
    return len(text.split())
